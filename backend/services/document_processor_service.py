"""Document Processor Service extracting text from PDF, Word, Excel, CSV, and Text files."""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any


class DocumentProcessorService:
    """Extracts text and metadata from supported file formats."""

    def extract_text(self, file_bytes: bytes, filename: str) -> tuple[str, int, int]:
        """Extract text from file bytes. Returns (text_content, page_count, text_length)."""
        ext = Path(filename).suffix.lower()

        if ext == ".txt":
            text_content = file_bytes.decode("utf-8", errors="replace")
            page_count = 1
        elif ext == ".csv":
            text_content = self._extract_csv(file_bytes)
            page_count = 1
        elif ext in {".pdf"}:
            text_content, page_count = self._extract_pdf(file_bytes)
        elif ext in {".docx", ".doc"}:
            text_content, page_count = self._extract_docx(file_bytes)
        elif ext in {".xlsx", ".xls"}:
            text_content, page_count = self._extract_excel(file_bytes)
        else:
            text_content = file_bytes.decode("utf-8", errors="replace")
            page_count = 1

        return text_content, page_count, len(text_content)

    def _extract_csv(self, file_bytes: bytes) -> str:
        text_stream = io.StringIO(file_bytes.decode("utf-8", errors="replace"))
        reader = csv.reader(text_stream)
        lines = []
        for row in reader:
            lines.append(", ".join(row))
        return "\n".join(lines)

    def _extract_pdf(self, file_bytes: bytes) -> tuple[str, int]:
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(file_bytes))
            pages_text = []
            for page in reader.pages:
                txt = page.extract_text() or ""
                pages_text.append(txt)
            full_text = "\n\n".join(pages_text)
            page_count = len(reader.pages) or 1
            return full_text if full_text.strip() else "PDF Document", page_count
        except Exception:
            text_fallback = file_bytes.decode("utf-8", errors="replace")
            return text_fallback, 1

    def _extract_docx(self, file_bytes: bytes) -> tuple[str, int]:
        try:
            import docx

            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text]
            full_text = "\n".join(paragraphs)
            return full_text if full_text.strip() else "Word Document", 1
        except Exception:
            text_fallback = file_bytes.decode("utf-8", errors="replace")
            return text_fallback, 1

    def _extract_excel(self, file_bytes: bytes) -> tuple[str, int]:
        try:
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
            sheet_texts = []
            for sheet in wb.worksheets:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(v) for v in row if v is not None]
                    if row_vals:
                        rows.append(", ".join(row_vals))
                sheet_texts.append("\n".join(rows))
            full_text = "\n\n".join(sheet_texts)
            return full_text if full_text.strip() else "Excel Document", len(wb.worksheets)
        except Exception:
            text_fallback = file_bytes.decode("utf-8", errors="replace")
            return text_fallback, 1
